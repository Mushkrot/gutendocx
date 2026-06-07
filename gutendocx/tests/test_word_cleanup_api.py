from docx import Document
from docx.enum.text import WD_BREAK

from gutendocx.web import server


def test_config_returns_word_cleanup_defaults(tmp_path):
    config_path = tmp_path / "config.yaml"

    data = server.get_config(config_path=str(config_path))

    assert data["ok"] is True
    assert data["word_cleanup"] == {
        "enabled": False,
        "patterns": [],
        "replacement": "^p",
    }


def test_word_cleanup_analyze_endpoint_returns_recommendations(tmp_path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Before")
    blank = doc.add_paragraph("")
    for _ in range(3):
        blank.add_run().add_break(WD_BREAK.LINE)
    doc.add_paragraph("After")
    doc.save(str(doc_path))

    data = server.word_cleanup_analyze(
        server.AnalyzeRequest(input=str(doc_path), config_path=str(tmp_path / "config.yaml")),
        None,
    )

    advisor = data["word_cleanup_advisor"]
    assert data["ok"] is True
    assert advisor["recommended_patterns"] == ["^p^l^l^l"]
    assert advisor["summary"]["recommended_gaps"] == 1
    assert advisor["summary"]["paragraphs_to_remove"] == 1
    assert advisor["summary"]["line_breaks_to_remove"] == 3
