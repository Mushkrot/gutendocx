import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from gutendocx.core import whole
from gutendocx.core.whole import _apply_para1_manual_line_break_style


def _normalizer():
    fn = getattr(whole, "_apply_body_style_normalization", None)
    if fn is None:
        pytest.xfail("body style normalizer is implemented in the next phase")
    return fn


def _set_keep_next(style, value):
    style.paragraph_format.keep_with_next = value


def _effective_keep_next(paragraph):
    direct = paragraph.paragraph_format.keep_with_next
    if direct is not None:
        return direct
    return paragraph.style.paragraph_format.keep_with_next


def _add_field_paragraph(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.text = " PAGE "
    instr_run._r.append(instr)
    paragraph.add_run("1")
    return paragraph


def test_bad_body_style_is_reassigned_to_safe_body_style():
    doc = Document()
    bad_style = doc.styles.add_style("Bad Body", WD_STYLE_TYPE.PARAGRAPH)
    bad_style.base_style = doc.styles["Normal"]
    _set_keep_next(bad_style, True)
    paragraph = doc.add_paragraph("Ordinary body text", style="Bad Body")

    result = _normalizer()(doc, {"body_style_normalization": {"enabled": True}}, 0)

    assert result["applied"] is True
    assert result["paragraphs_normalized"] == 1
    assert paragraph.style.name == "GD Body"
    assert _effective_keep_next(paragraph) is False
    assert doc.styles["GD Body"].paragraph_format.keep_with_next is False


def test_disabled_normalization_is_noop():
    doc = Document()
    bad_style = doc.styles.add_style("Bad Body", WD_STYLE_TYPE.PARAGRAPH)
    _set_keep_next(bad_style, True)
    paragraph = doc.add_paragraph("Ordinary body text", style="Bad Body")

    result = _normalizer()(doc, {"body_style_normalization": {"enabled": False}}, 0)

    assert result["applied"] is False
    assert result["enabled"] is False
    assert paragraph.style.name == "Bad Body"
    assert _effective_keep_next(paragraph) is True


def test_global_normal_style_is_not_modified_even_when_bad():
    doc = Document()
    _set_keep_next(doc.styles["Normal"], True)
    paragraph = doc.add_paragraph("Body paragraph inherited from bad Normal")

    result = _normalizer()(doc, {"body_style_normalization": {"enabled": True}}, 0)

    assert result["paragraphs_normalized"] == 1
    assert paragraph.style.name == "GD Body"
    assert _effective_keep_next(paragraph) is False
    assert doc.styles["Normal"].paragraph_format.keep_with_next is True


def test_headings_and_detected_heading_mappings_are_excluded():
    doc = Document()
    mapped_style = doc.styles.add_style("Para 08", WD_STYLE_TYPE.PARAGRAPH)
    _set_keep_next(mapped_style, True)
    heading = doc.add_paragraph("Built-in heading", style="Heading 1")
    mapped = doc.add_paragraph("Mapped heading", style="Para 08")
    config = {
        "body_style_normalization": {"enabled": True},
        "detected_style_mapping": {"Heading3": "Para 08"},
    }

    result = _normalizer()(doc, config, 0)

    assert result["paragraphs_normalized"] == 0
    assert heading.style.name == "Heading 1"
    assert mapped.style.name == "Para 08"
    assert mapped.style.paragraph_format.keep_with_next is True


def test_manual_line_break_body_still_becomes_para1_after_normalization():
    doc = Document()
    bad_style = doc.styles.add_style("Bad Body", WD_STYLE_TYPE.PARAGRAPH)
    _set_keep_next(bad_style, True)
    paragraph = doc.add_paragraph(style="Bad Body")
    paragraph.add_run("First line")
    paragraph.add_run().add_break(WD_BREAK.LINE)
    paragraph.add_run("Second line")

    _normalizer()(doc, {"body_style_normalization": {"enabled": True}}, 0)
    para1_result = _apply_para1_manual_line_break_style(doc, {}, 0)

    assert para1_result["paragraphs_modified"] == 1
    assert paragraph.style.name == "Para1"
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_toc_and_field_paragraphs_are_excluded():
    doc = Document()
    toc_style = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    _set_keep_next(toc_style, True)
    toc = doc.add_paragraph("Chapter 1\t3", style="TOC 1")
    field = _add_field_paragraph(doc)

    result = _normalizer()(doc, {"body_style_normalization": {"enabled": True}}, 0)

    assert result["paragraphs_normalized"] == 0
    assert toc.style.name == "TOC 1"
    assert field.style.name == "Normal"
