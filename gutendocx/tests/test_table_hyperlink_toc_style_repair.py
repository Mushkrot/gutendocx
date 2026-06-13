from zipfile import ZipFile

from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from gutendocx.core.toc import build_toc, repair_toc_result_runs, resolve_toc_style_overrides
from gutendocx.core.whole import apply_whole_document, restyle_body_nested_runs


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _w(tag):
    return f"{{{W_NS}}}{tag}"


def _body_config(tmp_path):
    return {
        "style_overrides": {
            "Body": {
                "font": "Aptos",
                "size_pt": 10,
                "bold": False,
                "italic": False,
            }
        },
        "output": {"dir": str(tmp_path)},
        "cover": {"normalize": {"map_runs_to_char_styles": False}},
    }


def _add_hyperlink(paragraph, text, url="https://example.com"):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return rid


def _run_style_ok(run_el, family="Aptos", half_points="20"):
    rpr = run_el.find(_w("rPr"))
    assert rpr is not None
    rfonts = rpr.find(_w("rFonts"))
    assert rfonts is not None
    assert rfonts.get(_w("ascii")) == family
    assert rfonts.get(_w("hAnsi")) == family
    assert rfonts.get(_w("eastAsia")) == family
    assert rfonts.get(_w("cs")) == family
    assert rpr.find(_w("sz")).get(_w("val")) == half_points
    assert rpr.find(_w("szCs")).get(_w("val")) == half_points


def test_body_styles_table_cells_and_hyperlink_visible_runs_without_touching_field_codes(tmp_path):
    doc = Document()
    doc.add_paragraph("Body")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("Table text")
    hyperlink_paragraph = doc.add_paragraph("See ")
    rid = _add_hyperlink(hyperlink_paragraph, "linked text")

    field_paragraph = doc.add_paragraph()
    field_run = field_paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.text = " PAGE "
    field_run._r.append(instr)

    input_path = tmp_path / "input.docx"
    doc.save(input_path)

    result = apply_whole_document(str(input_path), _body_config(tmp_path))
    output_path = result["output_path"]

    with ZipFile(output_path) as zf:
        xml = zf.read("word/document.xml")
    from lxml import etree

    root = etree.fromstring(xml)
    table_run = root.xpath(".//w:tbl//w:r[w:t]", namespaces={"w": W_NS})[0]
    _run_style_ok(table_run)

    hyperlink = root.xpath(".//w:hyperlink", namespaces={"w": W_NS})[0]
    assert hyperlink.get(qn("r:id")) == rid
    hyperlink_run = hyperlink.xpath(".//w:r[w:t]", namespaces={"w": W_NS})[0]
    _run_style_ok(hyperlink_run)

    instr_run = root.xpath(".//w:r[w:instrText]", namespaces={"w": W_NS})[0]
    assert instr_run.find(_w("rPr")) is None


def test_body_nested_runs_can_be_repaired_after_round_trip(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("Table text")
    hyperlink_paragraph = doc.add_paragraph("See ")
    _add_hyperlink(hyperlink_paragraph, "linked text")
    path = tmp_path / "repair.docx"
    doc.save(path)

    restyle_body_nested_runs(str(path), _body_config(tmp_path))

    with ZipFile(path) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    table_run = root.xpath(".//w:tbl//w:r[w:t]", namespaces={"w": W_NS})[0]
    hyperlink_run = root.xpath(".//w:hyperlink//w:r[w:t]", namespaces={"w": W_NS})[0]
    _run_style_ok(table_run)
    _run_style_ok(hyperlink_run)


def test_resolve_toc_style_modes():
    config = {
        "style_overrides": {
            "Body": {"font": "Aptos", "size_pt": 10},
            "Headings": {"font": "Georgia", "size_pt": 14},
            "Heading2": {"font": "Georgia", "size_pt": 12},
            "TOC": {"mode": "same_heading"},
        }
    }
    resolved = resolve_toc_style_overrides(config)
    assert resolved[1]["font_family"] == "Georgia"
    assert resolved[1]["size_pt"] == 14
    assert resolved[2]["size_pt"] == 12

    config["style_overrides"]["TOC"] = {
        "mode": "custom",
        "family": "Courier New",
        "size_pt": 9,
        "bold": False,
        "italic": True,
    }
    resolved = resolve_toc_style_overrides(config)
    assert resolved[1] == {
        "font_family": "Courier New",
        "size_pt": 9,
        "bold": False,
        "italic": True,
    }
    assert resolved[2] == resolved[1]


def test_toc_result_runs_inside_sdt_are_repaired_without_removing_fields():
    doc = Document()
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    sdt.append(content)

    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(qn("w:val"), "TOC1")
    ppr.append(pstyle)
    p.append(ppr)

    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    p.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\o "1-3" \\h '
    instr_run.append(instr)
    p.append(instr_run)

    text_run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "Chapter 1"
    text_run.append(text_node)
    p.append(text_run)

    content.append(p)
    doc._element.body.append(sdt)

    result = repair_toc_result_runs(
        doc,
        {"style_overrides": {"TOC": {"mode": "custom", "family": "Georgia", "size_pt": 9}}},
    )

    assert result["runs_modified"] == 1
    _run_style_ok(text_run, family="Georgia", half_points="18")
    assert begin_run.find(_w("rPr")) is None
    assert instr_run.find(_w("rPr")) is None
    assert p.find(".//" + _w("instrText")).text.startswith(" TOC")


def test_build_toc_preserves_existing_sdt_toc(tmp_path):
    doc = Document()
    doc.add_paragraph("Chapter 1").style = doc.styles["Heading 1"]
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    p = OxmlElement("w:p")
    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\o "1-3" \\h '
    instr_run.append(instr)
    p.append(instr_run)
    text_run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "Existing entry"
    text_run.append(text_node)
    p.append(text_run)
    content.append(p)
    doc._element.body.append(sdt)
    path = tmp_path / "toc.docx"
    doc.save(path)

    result = build_toc(
        str(path),
        {
            "output": {"dir": str(tmp_path), "versioning": True},
            "style_overrides": {"TOC": {"mode": "custom", "family": "Georgia", "size_pt": 9}},
        },
    )

    assert result["toc"]["inserted"]["inserted_at_end"] is False
    assert result["toc"]["removed"]["preserved_existing"] is True
    with ZipFile(result["output_path"]) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    assert len(root.xpath(".//w:sdt", namespaces={"w": W_NS})) == 1
    assert "Existing entry" in "".join(root.xpath(".//w:t/text()", namespaces={"w": W_NS}))
