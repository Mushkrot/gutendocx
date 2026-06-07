import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from gutendocx.core.word_cleanup import (
    analyze_word_cleanup,
    apply_word_cleanup,
    parse_word_cleanup_patterns,
)
from gutendocx.core.whole import _apply_para1_manual_line_break_style


def _cleanup_config(patterns):
    return {
        "word_cleanup": {
            "enabled": True,
            "patterns": patterns,
            "replacement": "^p",
        }
    }


def test_parse_word_cleanup_patterns_accepts_supported_codes_and_separators():
    assert parse_word_cleanup_patterns("^p, ^p^p; ^p^l^l^l^l\n^l") == [
        "^p",
        "^p^p",
        "^p^l^l^l^l",
        "^l",
    ]


def test_parse_word_cleanup_patterns_rejects_unknown_codes():
    with pytest.raises(ValueError, match="Use only \\^p and \\^l"):
        parse_word_cleanup_patterns("^p^x")


def test_cleanup_disabled_is_noop():
    doc = Document()
    doc.add_paragraph("Before")
    blank = doc.add_paragraph("")
    doc.add_paragraph("After")

    result = apply_word_cleanup(doc, {"word_cleanup": {"enabled": False, "patterns": ["^p"]}}, 0)

    assert result["applied"] is False
    assert len(doc.paragraphs) == 3
    assert doc.paragraphs[1].text == blank.text == ""


def test_cleanup_single_empty_body_paragraph_without_touching_text_boundaries():
    doc = Document()
    before = doc.add_paragraph("Before")
    doc.add_paragraph("")
    after = doc.add_paragraph("After")

    result = apply_word_cleanup(doc, _cleanup_config(["^p"]), 0)

    assert result["applied"] is True
    assert result["gaps_modified"] == 1
    assert result["paragraphs_removed"] == 1
    assert [p.text for p in doc.paragraphs] == ["Before", "After"]
    assert before.text == "Before"
    assert after.text == "After"


def test_cleanup_collapses_consecutive_empty_paragraphs():
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("After")

    result = apply_word_cleanup(doc, _cleanup_config(["^p^p"]), 0)

    assert result["gaps_modified"] == 1
    assert result["paragraphs_removed"] == 2
    assert [p.text for p in doc.paragraphs] == ["Before", "After"]


def test_cleanup_removes_empty_paragraph_with_four_manual_line_breaks():
    doc = Document()
    doc.add_paragraph("Before")
    blank = doc.add_paragraph("")
    for _ in range(4):
        blank.add_run().add_break(WD_BREAK.LINE)
    doc.add_paragraph("After")

    result = apply_word_cleanup(doc, _cleanup_config(["^p^l^l^l^l"]), 0)

    assert result["gaps_modified"] == 1
    assert result["paragraphs_removed"] == 1
    assert result["line_breaks_removed"] == 4
    assert [p.text for p in doc.paragraphs] == ["Before", "After"]


def test_nonempty_manual_line_break_paragraph_survives_for_para1_rule():
    doc = Document()
    p = doc.add_paragraph("")
    p.add_run("First")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("Second")

    result = apply_word_cleanup(doc, _cleanup_config(["^p^l"]), 0)
    para1 = _apply_para1_manual_line_break_style(doc, {}, 0)

    assert result["applied"] is False
    assert len(doc.paragraphs) == 1
    assert para1["paragraphs_modified"] == 1
    assert p.style.name == "Para1"


def test_cleanup_skips_heading_and_section_break_paragraphs():
    doc = Document()
    doc.add_paragraph("Before")
    heading_blank = doc.add_paragraph("", style="Heading 1")
    section_blank = doc.add_paragraph("")
    ppr = section_blank._element.get_or_add_pPr()
    ppr.append(OxmlElement("w:sectPr"))
    doc.add_paragraph("After")

    result = apply_word_cleanup(doc, _cleanup_config(["^p"]), 0)

    assert result["applied"] is False
    assert [p.text for p in doc.paragraphs] == ["Before", "", "", "After"]
    assert doc.paragraphs[1].style.name == heading_blank.style.name
    assert doc.paragraphs[2].text == section_blank.text


def test_analyze_word_cleanup_recommends_safe_patterns():
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_paragraph("")
    doc.add_paragraph("")
    blank_with_breaks = doc.add_paragraph("")
    for _ in range(4):
        blank_with_breaks.add_run().add_break(WD_BREAK.LINE)
    doc.add_paragraph("After")
    nonempty_break = doc.add_paragraph("")
    nonempty_break.add_run("Keep")
    nonempty_break.add_run().add_break(WD_BREAK.LINE)
    nonempty_break.add_run("This")

    advisor = analyze_word_cleanup(doc, {}, 0)

    assert advisor["recommended_patterns"] == ["^p^p", "^p^l^l^l^l"]
    assert advisor["summary"]["recommended_gaps"] == 1
    assert advisor["summary"]["paragraphs_to_remove"] == 3
    assert advisor["summary"]["line_breaks_to_remove"] == 4
    assert "Text paragraphs" in advisor["report"]
    assert [p.text for p in doc.paragraphs] == ["Before", "", "", "\n\n\n\n", "After", "Keep\nThis"]
