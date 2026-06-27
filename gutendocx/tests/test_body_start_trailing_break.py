import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from gutendocx.core.whole import _apply_body_style_overrides, _compute_body_start_index


def _add_page_break(paragraph):
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_section_break(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:sectPr"))


def test_single_trailing_page_break_is_ignored_as_body_boundary():
    doc = Document()
    doc.add_paragraph("Body text")
    trailing = doc.add_paragraph()
    _add_page_break(trailing)

    assert _compute_body_start_index(doc) == 0


def test_single_trailing_section_break_is_ignored_as_body_boundary():
    doc = Document()
    doc.add_paragraph("Body text")
    trailing = doc.add_paragraph()
    _add_section_break(trailing)

    assert _compute_body_start_index(doc) == 0


def test_single_nonfinal_break_still_starts_body_after_break():
    doc = Document()
    cover = doc.add_paragraph("Cover")
    _add_page_break(cover)
    doc.add_paragraph("Body text")

    assert _compute_body_start_index(doc) == 1


def test_two_breaks_still_start_body_after_second_break():
    doc = Document()
    cover = doc.add_paragraph("Cover")
    _add_page_break(cover)
    blank = doc.add_paragraph()
    _add_page_break(blank)
    doc.add_paragraph("Body text")

    assert _compute_body_start_index(doc) == 2


def test_line_spacing_still_applies_to_normal_whole_body_document():
    doc = Document()
    paragraph = doc.add_paragraph("Body text")
    config = {"style_overrides": {"Body": {"line_spacing": 1.08}}}

    result = _apply_body_style_overrides(doc, config, 0)

    assert result["changes"]["line_spacing"] == 1.08
    assert result["line_spacing_skipped_for_trailing_break_boundary"] == 0
    assert paragraph.paragraph_format.line_spacing == pytest.approx(1.08, abs=0.001)


def test_single_trailing_break_guard_skips_only_line_spacing_override():
    doc = Document()
    paragraph = doc.add_paragraph("Body text")
    trailing = doc.add_paragraph()
    _add_page_break(trailing)
    body_start = _compute_body_start_index(doc)
    config = {
        "style_overrides": {
            "Body": {
                "align": "justify",
                "line_spacing": 1.08,
            }
        }
    }

    result = _apply_body_style_overrides(doc, config, body_start)

    assert body_start == 0
    assert result["changes"] == {"alignment": "justify"}
    assert result["line_spacing_skipped_for_trailing_break_boundary"] == 2
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.line_spacing is None
