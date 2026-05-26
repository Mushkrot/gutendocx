from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

from gutendocx.core.whole import _apply_para1_manual_line_break_style


def test_manual_line_break_body_paragraph_gets_left_aligned_para1():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("First line")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("Second line")

    result = _apply_para1_manual_line_break_style(doc, {}, 0)

    assert result["applied"] is True
    assert result["paragraphs_modified"] == 1
    assert p.style.name == "Para1"
    assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert doc.styles["Para1"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_body_paragraph_without_manual_line_break_stays_normal():
    doc = Document()
    p = doc.add_paragraph("Plain body paragraph")

    result = _apply_para1_manual_line_break_style(doc, {}, 0)

    assert result["applied"] is False
    assert result["paragraphs_modified"] == 0
    assert p.style.name == "Normal"


def test_heading_with_manual_line_break_is_not_converted_to_para1():
    doc = Document()
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Chapter")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("One")

    result = _apply_para1_manual_line_break_style(doc, {}, 0)

    assert result["applied"] is False
    assert result["paragraphs_modified"] == 0
    assert p.style.name == "Heading 1"


def test_detected_heading_mapping_with_manual_line_break_is_not_converted_to_para1():
    doc = Document()
    doc.styles.add_style("Para 08", WD_STYLE_TYPE.PARAGRAPH)
    p = doc.add_paragraph(style="Para 08")
    p.add_run("Mapped chapter")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("Subtitle")
    config = {"detected_style_mapping": {"Heading3": "Para 08"}}

    result = _apply_para1_manual_line_break_style(doc, config, 0)

    assert result["applied"] is False
    assert result["paragraphs_modified"] == 0
    assert p.style.name == "Para 08"
