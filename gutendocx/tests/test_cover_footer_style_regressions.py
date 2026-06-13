from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from gutendocx.core.cover import apply_cover_styles
from gutendocx.core.layout import apply_footer_styles, restyle_footer_page_numbers


def _cover_config():
    return {
        "cover": {
            "styles": {
                "title": {
                    "name": "Cover Title",
                    "font": {"family": "Algerian", "size_pt": 24, "bold": False, "all_caps": False, "align": "center"},
                },
                "subtitle": {
                    "name": "Cover Subtitle",
                    "font": {"family": "Castellar", "size_pt": 18, "bold": False, "all_caps": False, "align": "center"},
                },
                "author": {
                    "name": "Cover Author",
                    "font": {"family": "Castellar", "size_pt": 18, "bold": False, "all_caps": False, "align": "center"},
                },
            },
            "normalize": {},
        }
    }


def _assert_no_direct_run_formatting(paragraph):
    for run in paragraph.runs:
        assert run._r.rPr is None


def test_cover_author_override_clears_direct_run_formatting():
    doc = Document()
    doc.add_paragraph("Title")
    author = doc.add_paragraph()
    run = author.add_run("Author")
    run.font.size = Pt(24)
    run.font.bold = True

    apply_cover_styles(
        doc,
        {"cover_paragraph_indices": [0, 1], "assignments": {0: "Cover Title", 1: "Cover Author"}},
        _cover_config(),
    )

    assert author.style.name == "Cover Author"
    _assert_no_direct_run_formatting(author)


def test_cover_subtitle_override_clears_direct_run_formatting():
    doc = Document()
    doc.add_paragraph("Title")
    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Subtitle")
    run.font.size = Pt(24)
    run.font.bold = True

    apply_cover_styles(
        doc,
        {"cover_paragraph_indices": [0, 1], "assignments": {0: "Cover Title", 1: "Cover Subtitle"}},
        _cover_config(),
    )

    assert subtitle.style.name == "Cover Subtitle"
    _assert_no_direct_run_formatting(subtitle)


def _footer_config():
    return {
        "style_overrides": {
            "Footer": {
                "font_family": "Times New Roman",
                "size_pt": 11,
                "bold": False,
                "italic": False,
            }
        }
    }


def _page_result_run(doc):
    footer = doc.sections[-1].footer
    for paragraph in footer.paragraphs:
        for run in paragraph.runs:
            if run.text == "1":
                return run
    raise AssertionError("PAGE result run not found")


def _assert_page_run_style(run):
    rpr = run._r.rPr
    assert rpr is not None
    rfonts = rpr.find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:ascii")) == "Times New Roman"
    assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
    assert rfonts.get(qn("w:eastAsia")) == "Times New Roman"
    assert rfonts.get(qn("w:cs")) == "Times New Roman"
    assert rpr.find(qn("w:sz")).get(qn("w:val")) == "22"
    assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "22"
    assert rpr.find(qn("w:b")).get(qn("w:val")) in {"0", "false"}
    assert rpr.find(qn("w:bCs")).get(qn("w:val")) == "0"
    assert rpr.find(qn("w:i")).get(qn("w:val")) in {"0", "false"}
    assert rpr.find(qn("w:iCs")).get(qn("w:val")) == "0"


def test_footer_page_field_has_explicit_result_run_style():
    doc = Document()
    doc.add_paragraph("Body")

    apply_footer_styles(doc, _footer_config())

    _assert_page_run_style(_page_result_run(doc))


def test_footer_page_style_can_be_repaired_after_round_trip(tmp_path):
    doc = Document()
    doc.add_paragraph("Body")
    apply_footer_styles(doc, _footer_config())
    path = tmp_path / "footer.docx"
    doc.save(path)

    stripped = Document(path)
    for paragraph in stripped.sections[-1].footer.paragraphs:
        for run in paragraph.runs:
            rpr = run._r.rPr
            if rpr is not None:
                run._r.remove(rpr)
    stripped.save(path)

    restyle_footer_page_numbers(str(path), _footer_config())

    repaired = Document(path)
    _assert_page_run_style(_page_result_run(repaired))
