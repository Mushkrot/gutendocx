from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

from gutendocx.core.whole import _compute_body_start_index


def _add_page_break(paragraph):
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_cover_paragraph(doc, text="Cover"):
    try:
        style = doc.styles["Cover Title"]
    except KeyError:
        style = doc.styles.add_style("Cover Title", WD_STYLE_TYPE.PARAGRAPH)
    return doc.add_paragraph(text, style=style)


def test_substantive_text_after_cover_is_recovered_before_second_break():
    doc = Document()
    cover = _add_cover_paragraph(doc)
    _add_page_break(cover)
    doc.add_paragraph("First body section " * 20)
    internal_break = doc.add_paragraph("End of first body section")
    _add_page_break(internal_break)
    doc.add_paragraph("Later body section")

    assert _compute_body_start_index(doc) == 1


def test_body_after_cover_is_recovered_when_cover_has_no_explicit_break():
    doc = Document()
    _add_cover_paragraph(doc)
    doc.add_paragraph()
    doc.add_paragraph("Biography " * 30)
    first_internal_break = doc.add_paragraph("End of biography")
    _add_page_break(first_internal_break)
    doc.add_paragraph("First poem " * 20)
    second_internal_break = doc.add_paragraph("End of first poem")
    _add_page_break(second_internal_break)
    doc.add_paragraph("Later poem")

    assert _compute_body_start_index(doc) == 2


def test_real_blank_interval_preserves_second_break_boundary():
    doc = Document()
    cover = _add_cover_paragraph(doc)
    _add_page_break(cover)
    blank = doc.add_paragraph()
    _add_page_break(blank)
    doc.add_paragraph("Body text")

    assert _compute_body_start_index(doc) == 2


def test_document_without_recognized_cover_style_preserves_legacy_boundary():
    doc = Document()
    first = doc.add_paragraph("Front matter " * 20)
    _add_page_break(first)
    second = doc.add_paragraph("More front matter " * 20)
    _add_page_break(second)
    doc.add_paragraph("Body text")

    assert _compute_body_start_index(doc) == 2


def test_short_decorative_marker_does_not_move_boundary():
    doc = Document()
    cover = _add_cover_paragraph(doc)
    _add_page_break(cover)
    doc.add_paragraph("***")
    blank = doc.add_paragraph()
    _add_page_break(blank)
    doc.add_paragraph("Body text")

    assert _compute_body_start_index(doc) == 3
